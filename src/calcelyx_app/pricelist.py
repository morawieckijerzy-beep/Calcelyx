"""Import, validation and export of editable construction pricelists.

Prices shipped with the application are demonstrative examples only.  A project
should keep its own imported snapshot instead of reading a changing workbook on
every opening.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path
import json
import math
import re
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .i18n import tr


FIELDS = (
    "code", "category", "name", "unit", "labor_net", "material_net",
    "vat_rate", "waste_pct", "min_charge", "active", "notes",
)
REQUIRED_FIELDS = FIELDS[:-1]
ALLOWED_UNITS = {
    "m²", "mb", "szt.", "kpl.", "pkt", "godz.", "dzień", "ryczałt",
    "lm", "pcs", "set", "point", "hour", "day", "lump sum",
}
CODE_PATTERN = re.compile(r"^[A-Z0-9][A-Z0-9_-]{1,39}$")
DEMO_NOTICE = (
    "CENNIK DEMONSTRACYJNY — kwoty są wyłącznie przykładowe i nie są cenami "
    "rynkowymi. Przed użyciem zastąp je własnymi stawkami."
)


@dataclass(frozen=True, slots=True)
class PriceItem:
    code: str
    category: str
    name: str
    unit: str
    labor_net: float
    material_net: float
    vat_rate: float
    waste_pct: float
    min_charge: float
    active: bool
    notes: str = ""


@dataclass(frozen=True, slots=True)
class ValidationIssue:
    row: int
    field: str
    message: str

    def __str__(self) -> str:
        where = f"wiersz {self.row}" if self.row else "arkusz"
        return tr(f"{where}, pole {self.field}: {self.message}")


class PricelistValidationError(ValueError):
    """Raised when a workbook contains one or more invalid rows."""

    def __init__(self, issues: Sequence[ValidationIssue]):
        self.issues = list(issues)
        super().__init__(tr("Nieprawidłowy cennik:\n") + "\n".join(map(str, self.issues)))


@dataclass(frozen=True, slots=True)
class PricelistImportResult:
    items: list[PriceItem]
    settings: dict[str, Any]


def default_data_path() -> Path:
    """Return the repository/bundle path of the demonstrative JSON catalogue."""
    return Path(__file__).resolve().parents[2] / "data" / "cennik_przykladowy.json"


def load_demo_items(path: str | Path | None = None) -> list[PriceItem]:
    source = Path(path) if path else default_data_path()
    raw = json.loads(source.read_text(encoding="utf-8"))
    return [PriceItem(**row) for row in raw["items"]]


def create_template_xlsx(
    path: str | Path,
    items: Iterable[PriceItem] | None = None,
    settings: Mapping[str, Any] | None = None,
    language: str = "pl",
) -> Path:
    """Create a styled, editable XLSX template and return its path."""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    currency = str((settings or {}).get("currency", "PLN")).strip().upper() or "PLN"
    workbook = Workbook()
    jobs = workbook.active
    jobs.title = "Work" if language == "en" else "Roboty"
    jobs.append(FIELDS)
    for item in (list(items) if items is not None else load_demo_items()):
        jobs.append([getattr(item, field) for field in FIELDS])
    _style_jobs_sheet(jobs, language, currency)

    settings_sheet = workbook.create_sheet("Settings" if language == "en" else "Ustawienia")
    settings_sheet.append(("key", "value", "description") if language == "en" else ("klucz", "wartość", "opis"))
    defaults = {
        "pricelist_name": "Example price list" if language == "en" else "Cennik demonstracyjny",
        "version": "1.0",
        "currency": "PLN",
        "default_vat_rate": 23,
        "valid_from": "",
        "demo_notice": "EXAMPLE PRICE LIST — all amounts are illustrative and are not market prices. Replace them with your own rates before use." if language == "en" else DEMO_NOTICE,
    }
    defaults.update(dict(settings or {}))
    descriptions = ({
        "pricelist_name": "Name displayed in the application",
        "version": "Price-list version stored with the estimate",
        "currency": "Currency code",
        "default_vat_rate": "Default VAT rate as a percentage",
        "valid_from": "Effective date, e.g. 2026-01-01",
        "demo_notice": "User notice",
    } if language == "en" else {
        "pricelist_name": "Nazwa widoczna w aplikacji",
        "version": "Wersja cennika zapisywana z wyceną",
        "currency": "Kod waluty",
        "default_vat_rate": "Domyślna stawka VAT w procentach",
        "valid_from": "Data obowiązywania, np. 2026-01-01",
        "demo_notice": "Ostrzeżenie dla użytkownika",
    })
    for key, value in defaults.items():
        settings_sheet.append((key, value, descriptions.get(key, "Custom setting" if language == "en" else "Ustawienie własne")))
    _style_table_sheet(settings_sheet, widths=(24, 54, 48))

    instructions = workbook.create_sheet("Instructions" if language == "en" else "Instrukcja")
    instructions.append(("How to use this price list" if language == "en" else "Jak korzystać z cennika",))
    lines = ([
        "EXAMPLE PRICE LIST — all amounts are illustrative and are not market prices. Replace them with your own rates before use.",
        "Do not change the column headers on the Work sheet.",
        "code must be unique and contain only uppercase letters, digits, _ or -.",
        "Allowed units: " + ", ".join(sorted(ALLOWED_UNITS)),
        "labor_net, material_net and min_charge are net amounts and cannot be negative.",
        "Enter vat_rate and waste_pct as numbers from 0 to 100, e.g. 23 and 10.",
        "active: enter YES/NO, TRUE/FALSE or 1/0.",
        "Do not delete the Work or Settings sheets.",
    ] if language == "en" else [
        DEMO_NOTICE,
        "Nie zmieniaj nazw nagłówków w arkuszu Roboty.",
        "code musi być unikalny i zawierać tylko wielkie litery, cyfry, _ lub -.",
        "Dozwolone jednostki: " + ", ".join(sorted(ALLOWED_UNITS)),
        "labor_net, material_net i min_charge to kwoty netto; nie mogą być ujemne.",
        "vat_rate i waste_pct wpisuj jako liczbę od 0 do 100, np. 23 i 10.",
        "active: wpisz TAK/NIE, TRUE/FALSE albo 1/0.",
        "Nie usuwaj arkuszy Roboty ani Ustawienia.",
    ])
    for line in lines:
        instructions.append((line,))
    instructions.column_dimensions["A"].width = 115
    instructions["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    instructions["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    for cell in instructions["A"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    instructions.freeze_panes = "A2"

    workbook.save(target)
    return target


def export_pricelist_xlsx(
    path: str | Path,
    items: Iterable[PriceItem],
    settings: Mapping[str, Any] | None = None,
) -> Path:
    """Export items using the same stable format as the editable template."""
    return create_template_xlsx(path, items=items, settings=settings)


def import_pricelist_xlsx(path: str | Path) -> PricelistImportResult:
    """Read an XLSX file, collecting all errors before rejecting it."""
    workbook = load_workbook(Path(path), data_only=True, read_only=True)
    issues: list[ValidationIssue] = []
    jobs_name = "Roboty" if "Roboty" in workbook.sheetnames else "Work" if "Work" in workbook.sheetnames else None
    if jobs_name is None:
        raise PricelistValidationError([ValidationIssue(0, "Roboty", "brak wymaganego arkusza")])
    jobs = workbook[jobs_name]
    rows = jobs.iter_rows(values_only=True)
    try:
        header_values = next(rows)
    except StopIteration:
        raise PricelistValidationError([ValidationIssue(1, "nagłówki", "arkusz jest pusty")])

    headers = [str(value).strip() if value is not None else "" for value in header_values]
    missing = [field for field in REQUIRED_FIELDS if field not in headers]
    if missing:
        raise PricelistValidationError([
            ValidationIssue(1, field, "brak wymaganej kolumny") for field in missing
        ])
    indexes = {name: headers.index(name) for name in FIELDS if name in headers}
    parsed: list[PriceItem] = []
    seen: dict[str, int] = {}
    for row_number, values in enumerate(rows, start=2):
        if not any(value not in (None, "") for value in values):
            continue
        raw = {field: values[index] if index < len(values) else None for field, index in indexes.items()}
        raw.setdefault("notes", "")
        item, row_issues = _parse_row(raw, row_number)
        issues.extend(row_issues)
        if item:
            if item.code in seen:
                issues.append(ValidationIssue(
                    row_number, "code", f"duplikat kodu {item.code!r}; pierwszy raz w wierszu {seen[item.code]}"
                ))
            else:
                seen[item.code] = row_number
                parsed.append(item)

    settings = _read_settings(workbook, issues)
    if issues:
        raise PricelistValidationError(issues)
    return PricelistImportResult(parsed, settings)


def _parse_row(raw: Mapping[str, Any], row: int) -> tuple[PriceItem | None, list[ValidationIssue]]:
    issues: list[ValidationIssue] = []
    code = _text(raw.get("code"))
    category = _text(raw.get("category"))
    name = _text(raw.get("name"))
    unit = _text(raw.get("unit"))
    if not code:
        issues.append(ValidationIssue(row, "code", "wartość jest wymagana"))
    elif not CODE_PATTERN.fullmatch(code):
        issues.append(ValidationIssue(row, "code", "użyj 2–40 znaków: A-Z, 0-9, _ lub -"))
    for field, value in (("category", category), ("name", name)):
        if not value:
            issues.append(ValidationIssue(row, field, "wartość jest wymagana"))
    if unit not in ALLOWED_UNITS:
        issues.append(ValidationIssue(row, "unit", f"niedozwolona jednostka {unit!r}"))

    numbers: dict[str, float] = {}
    for field in ("labor_net", "material_net", "vat_rate", "waste_pct", "min_charge"):
        value = raw.get(field)
        try:
            number = float(value)
            if not math.isfinite(number):
                raise ValueError
            numbers[field] = number
        except (TypeError, ValueError):
            issues.append(ValidationIssue(row, field, "oczekiwano liczby"))
    for field in ("labor_net", "material_net", "min_charge"):
        if field in numbers and numbers[field] < 0:
            issues.append(ValidationIssue(row, field, "wartość nie może być ujemna"))
    for field in ("vat_rate", "waste_pct"):
        if field in numbers and not 0 <= numbers[field] <= 100:
            issues.append(ValidationIssue(row, field, "wartość musi być w zakresie 0–100"))
    try:
        active = _bool(raw.get("active"))
    except ValueError:
        active = False
        issues.append(ValidationIssue(row, "active", "użyj TAK/NIE, TRUE/FALSE albo 1/0"))
    if issues:
        return None, issues
    return PriceItem(
        code, category, name, unit, numbers["labor_net"], numbers["material_net"],
        numbers["vat_rate"], numbers["waste_pct"], numbers["min_charge"], active,
        _text(raw.get("notes")),
    ), []


def _read_settings(workbook: Any, issues: list[ValidationIssue]) -> dict[str, Any]:
    settings_name = "Ustawienia" if "Ustawienia" in workbook.sheetnames else "Settings" if "Settings" in workbook.sheetnames else None
    if settings_name is None:
        issues.append(ValidationIssue(0, "Ustawienia", "brak wymaganego arkusza"))
        return {}
    sheet = workbook[settings_name]
    result: dict[str, Any] = {}
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        key = _text(values[0] if values else None)
        if not key:
            continue
        if key in result:
            issues.append(ValidationIssue(row_number, "klucz", f"duplikat ustawienia {key!r}"))
        else:
            result[key] = values[1] if len(values) > 1 else None
    return result


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    normalized = _text(value).casefold()
    if normalized in {"tak", "true", "1", "yes"}:
        return True
    if normalized in {"nie", "false", "0", "no"}:
        return False
    raise ValueError("invalid boolean")


def _style_jobs_sheet(sheet: Any, language: str = "pl", currency: str = "PLN") -> None:
    _style_table_sheet(sheet, widths=(16, 27, 43, 12, 15, 15, 12, 13, 15, 11, 48))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        row[9].value = ("YES" if language == "en" else "TAK") if bool(row[9].value) else ("NO" if language == "en" else "NIE")
        for index in (4, 5, 8):
            row[index].number_format = f'"{currency}" #,##0.00'
        for index in (6, 7):
            row[index].number_format = '0.00"%"'


def _style_table_sheet(sheet: Any, widths: Sequence[int]) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def item_to_dict(item: PriceItem) -> dict[str, Any]:
    """Stable serialization helper for storing a pricelist snapshot in a project."""
    return asdict(item)
