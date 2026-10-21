from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from .i18n import get_language, tr


ACCENT = colors.HexColor("#2563EB")
INK = colors.HexColor("#172033")
MUTED = colors.HexColor("#667085")
PALE = colors.HexColor("#EFF6FF")
LINE = colors.HexColor("#D9E2F0")
SUPPORTED_CURRENCIES = ("PLN", "EUR", "USD", "GBP", "CHF", "CZK", "SEK", "NOK", "DKK")


def _register_fonts() -> tuple[str, str]:
    candidates = [
        (Path("C:/Windows/Fonts/segoeui.ttf"), Path("C:/Windows/Fonts/seguisb.ttf")),
        (Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/arialbd.ttf")),
    ]
    for regular, bold in candidates:
        if regular.exists() and bold.exists():
            try:
                pdfmetrics.registerFont(TTFont("AppRegular", str(regular)))
                pdfmetrics.registerFont(TTFont("AppBold", str(bold)))
                return "AppRegular", "AppBold"
            except Exception:
                pass
    return "Helvetica", "Helvetica-Bold"


def normalize_currency(currency: Any) -> str:
    code = str(currency or "PLN").strip().upper()
    return code if code in SUPPORTED_CURRENCIES else "PLN"


def money(value: Any, language: str | None = None, currency: str = "PLN") -> str:
    amount = Decimal(str(value or 0)).quantize(Decimal("0.01"))
    selected = language or get_language()
    code = normalize_currency(currency)
    if selected == "pl":
        suffix = "zł" if code == "PLN" else code
        return f"{amount:,.2f} {suffix}".replace(",", " ").replace(".", ",")
    return f"{code} {amount:,.2f}"


def excel_currency_format(currency: str) -> str:
    """Return a portable Excel number format with an explicit ISO currency code."""
    return f'"{normalize_currency(currency)}" #,##0.00'


def safe(value: Any) -> str:
    return escape(str(value or ""))


def _totals(lines: Iterable[dict[str, Any]], discount_pct: Decimal, markup_pct: Decimal) -> dict[str, Decimal]:
    net = sum((Decimal(str(line.get("net", 0))) for line in lines if line.get("included", True)), Decimal("0"))
    markup = (net * markup_pct / Decimal("100")).quantize(Decimal("0.01"))
    discount = ((net + markup) * discount_pct / Decimal("100")).quantize(Decimal("0.01"))
    after = net + markup - discount
    base_vat = sum(
        (
            Decimal(str(line.get("net", 0)))
            * Decimal(str(line.get("vat_rate", 23)))
            / Decimal("100")
            for line in lines
            if line.get("included", True)
        ),
        Decimal("0"),
    )
    factor = (Decimal("1") + markup_pct / Decimal("100")) * (
        Decimal("1") - discount_pct / Decimal("100")
    )
    vat = (base_vat * factor).quantize(Decimal("0.01"))
    return {"base": net, "markup": markup, "discount": discount, "net": after, "vat": vat, "gross": after + vat}


def export_offer_pdf(path: str | Path, project: dict[str, Any], lines: list[dict[str, Any]], language: str = "pl") -> Path:
    """Create a customer-facing, A4 offer PDF from plain dictionaries."""
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    regular, bold = _register_fonts()
    styles = getSampleStyleSheet()
    normal = ParagraphStyle("Normal", parent=styles["Normal"], fontName=regular, fontSize=9, leading=13, textColor=INK)
    small = ParagraphStyle("Small", parent=normal, fontSize=7.5, leading=10, textColor=MUTED)
    title = ParagraphStyle("Title", parent=normal, fontName=bold, fontSize=23, leading=27, textColor=INK)
    heading = ParagraphStyle("Heading", parent=normal, fontName=bold, fontSize=12, leading=16, textColor=INK, spaceBefore=7, spaceAfter=5)
    right = ParagraphStyle("Right", parent=normal, alignment=TA_RIGHT)
    t = lambda text: tr(text, language)
    currency = normalize_currency(project.get("currency"))
    formatted_money = lambda value: money(value, language, currency)
    number = lambda value: f"{Decimal(str(value or 0)):.2f}".replace(".", ",") if language == "pl" else f"{Decimal(str(value or 0)):.2f}"

    discount = Decimal(str(project.get("discount_pct", 0)))
    markup = Decimal(str(project.get("markup_pct", 0)))
    totals = _totals(lines, discount, markup)

    doc = SimpleDocTemplate(
        str(output),
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=f"{t('Oferta')} {project.get('number', '')}",
        author=project.get("company_name", "Calcelyx"),
    )
    story: list[Any] = []
    header = Table(
        [
            [Paragraph(safe(project.get("company_name", t("TWOJA FIRMA"))), heading), Paragraph(t("OFERTA WYCENY"), title)],
            [Paragraph(safe(project.get("company_details", t("Dane firmy do uzupełnienia"))), small), Paragraph(f"{t('Nr')} {safe(project.get('number', 'OF/001'))}<br/>{safe(project.get('date', date.today().isoformat()))}", right)],
        ],
        colWidths=[82 * mm, 96 * mm],
    )
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, -1), (-1, -1), 1.2, ACCENT), ("BOTTOMPADDING", (0, -1), (-1, -1), 9)]))
    story.extend([header, Spacer(1, 8 * mm)])

    info = Table(
        [
            [Paragraph(t("KLIENT"), small), Paragraph(t("INWESTYCJA"), small), Paragraph(t("WAŻNOŚĆ"), small)],
            [Paragraph(safe(project.get("client_name", "—")), normal), Paragraph(safe(project.get("address", "—")), normal), Paragraph(f"{safe(project.get('valid_days', 14))} {t('dni')}", normal)],
        ],
        colWidths=[58 * mm, 76 * mm, 44 * mm],
    )
    info.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), PALE), ("BOX", (0, 0), (-1, -1), 0.5, LINE), ("INNERGRID", (0, 0), (-1, -1), 0.5, LINE), ("PADDING", (0, 0), (-1, -1), 7)]))
    story.extend([info, Spacer(1, 8 * mm), Paragraph(t("Zakres i wycena"), heading)])

    table_data: list[list[Any]] = [[t("Lp."), t("Pomieszczenie / praca"), t("Ilość"), t("J.m."), t("Cena netto"), t("Wartość netto")]]
    for index, line in enumerate((x for x in lines if x.get("included", True)), 1):
        label = f"<b>{safe(line.get('name', ''))}</b>"
        if line.get("room"):
            label = f"{safe(line['room'])}<br/><font color='#667085'>{safe(line.get('name', ''))}</font>"
        if line.get("formula"):
            label += f"<br/><font size='7' color='#667085'>{safe(line['formula'])}</font>"
        table_data.append([
            str(index), Paragraph(label, normal), number(line.get("quantity", 0)), line.get("unit", ""), formatted_money(line.get("unit_price", 0)), formatted_money(line.get("net", 0)),
        ])
    detail = Table(table_data, repeatRows=1, colWidths=[9 * mm, 73 * mm, 20 * mm, 13 * mm, 29 * mm, 34 * mm])
    detail.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), INK), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), bold),
        ("FONTNAME", (0, 1), (-1, -1), regular), ("FONTSIZE", (0, 0), (-1, -1), 7.8), ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("GRID", (0, 0), (-1, -1), 0.35, LINE), ("PADDING", (0, 0), (-1, -1), 5),
    ]))
    story.extend([detail, Spacer(1, 6 * mm)])

    summary_rows = [[t("Suma pozycji"), formatted_money(totals["base"])]]
    if totals["markup"]:
        summary_rows.append([f"{t('Narzut')} {markup}%", formatted_money(totals["markup"])])
    if totals["discount"]:
        summary_rows.append([f"{t('Rabat')} {discount}%", f"− {formatted_money(totals['discount'])}"])
    summary_rows.extend([[t("Razem netto"), formatted_money(totals["net"])], ["VAT", formatted_money(totals["vat"])], [t("RAZEM BRUTTO"), formatted_money(totals["gross"])]])
    summary = Table(summary_rows, colWidths=[43 * mm, 40 * mm], hAlign="RIGHT")
    summary.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -2), regular), ("FONTNAME", (0, -1), (-1, -1), bold), ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"), ("LINEABOVE", (0, -1), (-1, -1), 1.2, ACCENT), ("BACKGROUND", (0, -1), (-1, -1), PALE), ("PADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(summary)

    assumptions = project.get("assumptions", "")
    if assumptions:
        story.extend([Spacer(1, 7 * mm), Paragraph(t("Założenia i uwagi"), heading), Paragraph(safe(assumptions).replace("\n", "<br/>"), normal)])
    excluded = [x for x in lines if not x.get("included", True)]
    if excluded:
        story.extend([Spacer(1, 5 * mm), Paragraph(t("Opcje nieuwzględnione w sumie"), heading)])
        for item in excluded:
            story.append(Paragraph(f"• {safe(item.get('name', ''))}: {formatted_money(item.get('net', 0))} {t('netto')}", normal))

    story.extend([
        Spacer(1, 12 * mm),
        Table([["____________________________", "____________________________"], [t("Wykonawca"), t("Akceptacja klienta")]], colWidths=[89 * mm, 89 * mm], style=TableStyle([("FONTNAME", (0, 0), (-1, -1), regular), ("TEXTCOLOR", (0, 1), (-1, 1), MUTED), ("FONTSIZE", (0, 0), (-1, -1), 8), ("ALIGN", (1, 0), (1, -1), "RIGHT")]))
    ])

    def footer(canvas: Any, document: Any) -> None:
        canvas.saveState()
        canvas.setFont(regular, 7)
        canvas.setFillColor(MUTED)
        canvas.drawString(16 * mm, 8 * mm, f"{t('Oferta')} {project.get('number', '')} • {t('wygenerowano w')} Calcelyx 1.0")
        canvas.drawRightString(194 * mm, 8 * mm, f"{t('Strona')} {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output


def export_estimate_xlsx(path: str | Path, project: dict[str, Any], lines: list[dict[str, Any]], language: str = "pl") -> Path:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    t = lambda text: tr(text, language)
    currency = normalize_currency(project.get("currency"))
    ws.title = t("Kosztorys")
    ws.append([t("Oferta"), project.get("number", ""), t("Klient"), project.get("client_name", ""), t("Adres"), project.get("address", ""), t("Waluta"), currency])
    ws.append([])
    headers = [t(value) for value in ["Lp.", "Pomieszczenie", "Kod", "Nazwa", "Ilość", "J.m.", "Robocizna netto/j.m.", "Materiał netto/j.m.", "Cena łączna netto/j.m.", "Wartość netto", "VAT %", "Wliczona", "Wzór / uwagi"]]
    ws.append(headers)
    for cell in ws[3]:
        cell.fill = PatternFill("solid", fgColor="172033")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    for i, line in enumerate(lines, 1):
        ws.append([i, line.get("room", ""), line.get("code", ""), line.get("name", ""), float(line.get("quantity", 0)), line.get("unit", ""), float(line.get("labor_net", line.get("unit_price", 0))), float(line.get("material_net", 0)), float(line.get("unit_price", 0)), float(line.get("net", 0)), float(line.get("vat_rate", 23)), t("TAK") if line.get("included", True) else t("NIE"), line.get("formula", "")])
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:M{ws.max_row}"
    widths = [7, 20, 15, 42, 12, 10, 18, 18, 19, 17, 9, 11, 38]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in ws.iter_rows(min_row=4, min_col=5, max_col=11):
        for cell in row:
            cell.number_format = '#,##0.00'
    for row in ws.iter_rows(min_row=4, min_col=7, max_col=10):
        for cell in row:
            cell.number_format = excel_currency_format(currency)
    wb.save(output)
    return output

