from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from calcelyx_app.app import App
from calcelyx_app.i18n import set_language, tr
from calcelyx_app.pricelist import create_template_xlsx, import_pricelist_xlsx, load_demo_items
from calcelyx_app.reporting import export_estimate_xlsx, export_offer_pdf, money


PROJECT = {
    "number": "QT/2026/001",
    "date": "2026-10-21",
    "client_name": "Test Client",
    "address": "Test Street 1",
    "valid_days": 14,
    "company_name": "Test Company",
    "company_details": "test@example.com",
    "markup_pct": 0,
    "discount_pct": 0,
    "assumptions": "Measurements must be verified on site.",
}
LINES = [{
    "room": "Living room", "code": "TEST_M2", "name": "Paint walls", "quantity": 10,
    "unit": "m²", "labor_net": 20, "material_net": 5, "unit_price": 25,
    "net": 250, "vat_rate": 23, "included": True, "formula": "10 × PLN 25.00",
}]


def test_english_is_default_for_new_projects() -> None:
    set_language("en")
    state = App._default_state()
    assert state["language"] == "en"
    assert state["project"]["name"] == "New estimate"
    assert state["project"]["number"].startswith("QT/")
    assert state["project"]["currency"] == "PLN"


def test_translation_handles_dynamic_messages_but_preserves_user_text() -> None:
    assert tr("Zapisano lokalnie • 12:30", "en") == "Saved locally • 12:30"
    assert tr("Data House", "en") == "Data House"
    assert tr("Ustawienia", "pl") == "Ustawienia"


def test_both_demo_catalogues_have_matching_codes_and_prices() -> None:
    root = Path(__file__).resolve().parents[1] / "data"
    polish = load_demo_items(root / "cennik_przykladowy.json")
    english = load_demo_items(root / "pricelist_example_en.json")
    assert len(polish) == len(english) == 51
    assert [(x.code, x.labor_net, x.material_net) for x in polish] == [
        (x.code, x.labor_net, x.material_net) for x in english
    ]


def test_english_price_list_round_trip(tmp_path: Path) -> None:
    root = Path(__file__).resolve().parents[1] / "data"
    items = load_demo_items(root / "pricelist_example_en.json")
    target = create_template_xlsx(tmp_path / "price_list.xlsx", items, language="en")
    workbook = load_workbook(target, read_only=True, data_only=True)
    assert workbook.sheetnames == ["Work", "Settings", "Instructions"]
    imported = import_pricelist_xlsx(target)
    assert len(imported.items) == 51
    assert imported.items[0].name == "Protect floors and furnishings"


def test_english_exports_use_english_labels(tmp_path: Path) -> None:
    pdf = export_offer_pdf(tmp_path / "quote.pdf", PROJECT, LINES, language="en")
    text = "\n".join(page.extract_text() or "" for page in PdfReader(pdf).pages)
    assert "PROJECT QUOTE" in text
    assert "GROSS TOTAL" in text

    xlsx = export_estimate_xlsx(tmp_path / "estimate.xlsx", PROJECT, LINES, language="en")
    workbook = load_workbook(xlsx, read_only=True, data_only=True)
    assert workbook.sheetnames == ["Estimate"]
    assert workbook["Estimate"]["A3"].value == "No."
    assert workbook["Estimate"]["B3"].value == "Room"


def test_currency_is_localized_and_exported(tmp_path: Path) -> None:
    assert money("1234.5", language="en", currency="EUR") == "EUR 1,234.50"
    assert money("1234.5", language="pl", currency="EUR") == "1 234,50 EUR"
    assert money("1234.5", language="pl", currency="PLN") == "1 234,50 zł"

    project = {**PROJECT, "currency": "EUR"}
    pdf = export_offer_pdf(tmp_path / "quote_eur.pdf", project, LINES, language="en")
    text = "\n".join(page.extract_text() or "" for page in PdfReader(pdf).pages)
    assert "EUR 250.00" in text

    xlsx = export_estimate_xlsx(tmp_path / "estimate_eur.xlsx", project, LINES, language="en")
    workbook = load_workbook(xlsx, data_only=False)
    sheet = workbook["Estimate"]
    assert sheet["G1"].value == "Currency"
    assert sheet["H1"].value == "EUR"
    assert "EUR" in sheet["J4"].number_format
