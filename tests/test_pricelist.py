from __future__ import annotations

from openpyxl import load_workbook
import pytest

from calcelyx_app.pricelist import (
    DEMO_NOTICE,
    PricelistValidationError,
    create_template_xlsx,
    export_pricelist_xlsx,
    import_pricelist_xlsx,
    load_demo_items,
)


def test_demo_catalogue_is_large_and_clearly_marked() -> None:
    items = load_demo_items()
    assert len(items) >= 45
    assert len({item.code for item in items}) == len(items)
    assert len({item.category for item in items}) >= 9
    assert all("DEMO" in item.notes for item in items)
    assert "nie są cenami rynkowymi" in DEMO_NOTICE


def test_xlsx_round_trip_preserves_items_and_settings(tmp_path) -> None:
    original = load_demo_items()
    target = tmp_path / "cennik.xlsx"
    export_pricelist_xlsx(
        target,
        original,
        settings={"pricelist_name": "Cennik testowy", "version": "2026.10"},
    )

    result = import_pricelist_xlsx(target)

    assert result.items == original
    assert result.settings["pricelist_name"] == "Cennik testowy"
    assert result.settings["version"] == "2026.10"
    workbook = load_workbook(target, read_only=True)
    assert workbook.sheetnames == ["Roboty", "Ustawienia", "Instrukcja"]


def test_template_uses_selected_currency_in_settings_and_price_formats(tmp_path) -> None:
    target = create_template_xlsx(
        tmp_path / "cennik_eur.xlsx",
        items=load_demo_items()[:1],
        settings={"currency": "EUR"},
    )
    result = import_pricelist_xlsx(target)
    workbook = load_workbook(target, data_only=False)
    assert result.settings["currency"] == "EUR"
    assert "EUR" in workbook["Roboty"]["E2"].number_format


def test_template_uses_demo_catalogue_by_default(tmp_path) -> None:
    target = create_template_xlsx(tmp_path / "szablon.xlsx")
    result = import_pricelist_xlsx(target)
    assert len(result.items) >= 45
    assert result.settings["demo_notice"] == DEMO_NOTICE


def test_import_reports_duplicate_bad_unit_and_bad_numbers_with_rows(tmp_path) -> None:
    target = create_template_xlsx(tmp_path / "bledny.xlsx", items=load_demo_items()[:2])
    workbook = load_workbook(target)
    sheet = workbook["Roboty"]
    sheet["A3"] = sheet["A2"].value  # duplicate code
    sheet["D2"] = "metr"             # invalid unit
    sheet["E2"] = "drogo"            # invalid number
    sheet["G3"] = 101                 # percentage out of range
    workbook.save(target)

    with pytest.raises(PricelistValidationError) as caught:
        import_pricelist_xlsx(target)

    issues = caught.value.issues
    assert any(issue.row == 2 and issue.field == "unit" for issue in issues)
    assert any(issue.row == 2 and issue.field == "labor_net" for issue in issues)
    assert any(issue.row == 3 and issue.field == "vat_rate" for issue in issues)
    # A row with any invalid field is not accepted as a valid candidate, but the
    # duplicate is still independently tested below to keep error semantics clear.


def test_import_reports_duplicate_codes(tmp_path) -> None:
    target = create_template_xlsx(tmp_path / "duplikat.xlsx", items=load_demo_items()[:2])
    workbook = load_workbook(target)
    sheet = workbook["Roboty"]
    sheet["A3"] = sheet["A2"].value
    workbook.save(target)

    with pytest.raises(PricelistValidationError) as caught:
        import_pricelist_xlsx(target)

    assert any(issue.row == 3 and issue.field == "code" and "duplikat" in issue.message
               for issue in caught.value.issues)


def test_import_rejects_missing_required_sheet(tmp_path) -> None:
    target = create_template_xlsx(tmp_path / "bez_robot.xlsx", items=[])
    workbook = load_workbook(target)
    del workbook["Roboty"]
    workbook.save(target)

    with pytest.raises(PricelistValidationError) as caught:
        import_pricelist_xlsx(target)

    assert caught.value.issues[0].field == "Roboty"
