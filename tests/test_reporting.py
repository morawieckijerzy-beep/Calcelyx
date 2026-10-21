from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from decimal import Decimal

from calcelyx_app.reporting import _totals, export_estimate_xlsx, export_offer_pdf


PROJECT = {
    "company_name": "Firma Testowa",
    "company_details": "NIP 000",
    "number": "OF/TEST/1",
    "date": "2026-10-21",
    "client_name": "Jan Kowalski",
    "address": "Warszawa",
    "valid_days": 14,
    "markup_pct": 5,
    "discount_pct": 2,
    "assumptions": "Wymiary zweryfikować na miejscu.",
}

LINES = [
    {"room": "Salon", "code": "MAL-01", "name": "Malowanie ścian", "quantity": "42.5", "unit": "m²", "unit_price": "20", "net": "850", "vat_rate": 8, "included": True, "formula": "42,5 m² × 20 zł"},
    {"room": "Salon", "code": "OPT-01", "name": "Opcja dodatkowa", "quantity": "1", "unit": "kpl.", "unit_price": "100", "net": "100", "vat_rate": 23, "included": False},
]


def test_pdf_export_is_readable(tmp_path: Path):
    target = export_offer_pdf(tmp_path / "oferta.pdf", PROJECT, LINES)
    reader = PdfReader(target)
    assert len(reader.pages) >= 1
    assert target.stat().st_size > 2_000


def test_xlsx_export_has_formula_and_rows(tmp_path: Path):
    target = export_estimate_xlsx(tmp_path / "kosztorys.xlsx", PROJECT, LINES)
    workbook = load_workbook(target, data_only=False)
    sheet = workbook["Kosztorys"]
    assert sheet.max_row == 5
    assert sheet["J4"].value == 850


def test_pdf_totals_apply_markup_and_discount_to_vat():
    totals = _totals(LINES, Decimal("2"), Decimal("5"))
    factor = Decimal("1.05") * Decimal("0.98")
    expected_vat = (Decimal("850") * Decimal("0.08") * factor).quantize(Decimal("0.01"))
    assert totals["vat"] == expected_vat

